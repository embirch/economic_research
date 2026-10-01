#!/usr/bin/env python3
"""Extract the sex rows of the Cetic.br TIC Domicilios 2025 generative-AI tables.

Input: the provider's published individuals table bundle
`ict_households_2025_individuals_tables_xlsx_v1.0.zip`, downloaded from
https://cetic.br/pt/arquivos/domicilios/2025/individuos/ (public aggregate tables;
no respondent microdata is read or required).

The script pins the exact file vintage by SHA-256, then reads the SEX rows of
sheets M1 (use), M2 (purposes among users) and M3 (declared reasons for not using)
from the proportion, total and sampling-error workbooks. It prints JSON and exits
non-zero if the archive hash, sheet names or row labels differ, so a later
download cannot be silently substituted.

Usage:
    python3 check_cetic.py /path/to/ict_households_2025_individuals_tables_xlsx_v1.0.zip

Requires: openpyxl.
"""

import hashlib
import io
import json
import sys
import zipfile

import openpyxl

EXPECTED_SHA256 = "9ca6638cc6a965d7d90de789cdc7f0e67de08918d4f4067c56575e9b621dd930"
EXPECTED_BYTES = 1027291

MEMBERS = {
    "proportion": "ict_households_2025_individuals_table_proportion_v1.0.xlsx",
    "total": "ict_households_2025_individuals_table_total_v1.0.xlsx",
    "margin_of_error": "ict_households_2025_individuals_table_sampling_error_v1.0.xlsx",
}

SHEETS = {
    "M1": "internet users by use of generative AI tools",
    "M2": "internet users who used generative AI, by purpose of use",
    "M3": "internet users who did not use generative AI, by declared reason",
}


def read_sheet(data, sheet):
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise SystemExit(f"sheet {sheet} missing; found {wb.sheetnames}")
    ws = wb[sheet]
    title = None
    base = None
    columns = None
    rows = {}
    for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        cells = ["" if c is None else c for c in row]
        if idx == 1:
            title = str(cells[0]).strip()
            continue
        if idx == 2:
            base = str(cells[0]).strip()
            continue
        if idx == 3:
            columns = [str(c).strip() for c in cells[2:] if str(c).strip()]
            continue
        if str(cells[0]).strip() == "SEX":
            label = str(cells[1]).strip()
            rows[label] = [c for c in cells[2 : 2 + len(columns)]]
        if idx > 60:
            break
    wb.close()
    if sorted(rows) != ["Female", "Male"]:
        raise SystemExit(f"expected Male/Female SEX rows in {sheet}, found {sorted(rows)}")
    return {"sheet_title": title, "base": base, "columns": columns, "sex_rows": rows}


def main(path):
    raw = open(path, "rb").read()
    digest = hashlib.sha256(raw).hexdigest()
    if digest != EXPECTED_SHA256 or len(raw) != EXPECTED_BYTES:
        raise SystemExit(
            "archive differs from the audited vintage: "
            f"{len(raw)} bytes, sha256 {digest}; expected {EXPECTED_BYTES} bytes, "
            f"sha256 {EXPECTED_SHA256}. Re-audit before reuse."
        )
    zf = zipfile.ZipFile(io.BytesIO(raw))
    out = {
        "source_id_prefix": "BR_CETIC",
        "sample_family": "BR_CETIC_TICDOM_2025",
        "provider": "Cetic.br / NIC.br",
        "survey": "TIC Domicilios 2025 (individuals tables, version 1.0)",
        "retrieved": "2026-10-01",
        "source_url": "https://cetic.br/media/microdados/986/ict_households_2025_individuals_tables_xlsx_v1.0.zip",
        "archive_sha256": digest,
        "archive_bytes": len(raw),
        "tables": {},
    }
    for sheet, description in SHEETS.items():
        out["tables"][sheet] = {"description": description, "measures": {}}
        for measure, member in MEMBERS.items():
            if member not in zf.namelist():
                raise SystemExit(f"missing workbook {member}")
            out["tables"][sheet]["measures"][measure] = read_sheet(zf.read(member), sheet)
    print(json.dumps(out, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    main(sys.argv[1])
